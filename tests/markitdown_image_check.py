"""Inside the markitdown image: one document of each format protocol routes
there, through the port protocol calls, with the network off.

    docker run --rm -i --network=none markitdown-service:ci python - < tests/markitdown_image_check.py

CI runs it (job `markitdown-image`) because the image is small enough to build
there, unlike marker's: a markitdown bump that breaks a reader then fails a
job instead of a conversion in protocol. Not a unittest module (the name does
not match test*.py) - it needs the image's libraries.

The documents are generated here rather than kept as fixtures: DOCX and EPUB
by hand with zipfile (each is a zip of a few XML parts), PPTX and XLSX with
python-pptx and openpyxl, which the image's extras already install.
"""
import os
import sys
import tempfile
import zipfile

import markitdown_tasks

MARK = "PROTOCOLDROIDCHECK"
tmp = tempfile.mkdtemp(prefix="md-check-")


def docx():
    path = os.path.join(tmp, "doc.docx")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/'
                   'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                   'relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                   f'<w:body><w:p><w:r><w:t>{MARK} docx</w:t></w:r></w:p></w:body></w:document>')
    return path


def epub():
    path = os.path.join(tmp, "book.epub")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>')
        z.writestr("content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
                   'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
                   '<dc:title>Check</dc:title><dc:identifier id="id">check</dc:identifier></metadata>'
                   '<manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="c1"/></spine></package>')
        z.writestr("c1.xhtml",
                   '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><body>'
                   f'<p>{MARK} epub</p></body></html>')
    return path


def pptx():
    from pptx import Presentation
    path = os.path.join(tmp, "deck.pptx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = f"{MARK} pptx"
    deck.save(path)
    return path


def xlsx():
    from openpyxl import Workbook
    path = os.path.join(tmp, "sheet.xlsx")
    book = Workbook()
    book.active.append(["name", "value"])
    book.active.append([f"{MARK} xlsx", 1])
    book.save(path)
    return path


def text(name, body):
    path = os.path.join(tmp, name)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(body)
    return path


failures, versions = [], {}
# Factories, called inside the try: a reader library missing from the image
# is then reported against its format instead of stopping the whole check.
cases = [
    ("docx", docx), ("epub", epub), ("pptx", pptx), ("xlsx", xlsx),
    ("html", lambda: text("page.html", f"<html><body><h1>{MARK} html</h1></body></html>")),
    ("csv", lambda: text("table.csv", f"name,value\n{MARK} csv,1\n")),
]
for kind, make in cases:
    try:
        path = make()
        with open(path, "rb") as fh:
            files, versions = markitdown_tasks.convert_bytes(fh.read(), os.path.basename(path))
        (name, body), = files.items()
        ok = f"{MARK} {kind}" in body.decode("utf-8")
        print(f"{'ok  ' if ok else 'FAIL'} {kind:5} -> {name} ({len(body)} bytes)")
        if not ok:
            failures.append(f"{kind}: marker text missing from {body[:200]!r}")
    except Exception as exc:  # noqa: BLE001 - report every format, then fail
        print(f"FAIL {kind:5} {type(exc).__name__}: {exc}")
        failures.append(kind)

try:
    markitdown_tasks.convert_bytes(b"x", "a.docx", {"page_range": "0"})
    failures.append("options were accepted")
    print("FAIL options accepted")
except markitdown_tasks.InvalidOption:
    print("ok   options refused")

print("versions:", versions)
if versions.get("markitdown", "absent") == "absent":
    failures.append("markitdown version not recorded")
if failures:
    print("FAILED:", "; ".join(failures))
    sys.exit(1)
print("all formats converted")
